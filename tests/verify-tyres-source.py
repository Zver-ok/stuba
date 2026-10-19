"""Check the published tyre data against the supplied workbook, without reusing the importer."""
import json
import re
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
sheet = openpyxl.load_workbook(root / 'assets/docs/tyres-specifications.xlsx', data_only=True).active
models = json.loads((root / 'tyres-data.json').read_text())['models']
headings = [cell.row for row in sheet for cell in row[:1] if isinstance(cell.value, str) and re.match(r'^\d+\.\s*Шина', cell.value)]
assert len(models) == len(headings) == 25
assert {model['sourceRow'] for model in models} == set(headings)
for model in models:
    row = model['sourceRow']
    assert re.sub(r'^\d+\.\s*Шина\s*', '', sheet.cell(row, 1).value).strip() == model['name']
    assert int(re.sub(r'\D', '', sheet.cell(row + 1, 1).value)) == model['price']
    actual = []
    row += 2
    while row <= sheet.max_row and sheet.cell(row, 2).value is not None:
        actual.append([str(sheet.cell(row, 1).value).strip(), str(sheet.cell(row, 2).value).strip()])
        row += 1
    assert actual == model['specs']['Характеристики']
assert {model['brand'] for model in models} == {'TRONMAX', 'ARMOUR'}
print('PASS: all 25 XLSX positions, exact names, prices and specification cells; no added products.')
